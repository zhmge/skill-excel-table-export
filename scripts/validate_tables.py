#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""validate_tables.py — Excel 配置表导表工具前置静态校验（只读）。

在进 Unity 执行 Tools/Export Tables 之前，先对 Excel 配置表做结构检查，
把表头缺失、类型不支持、表名冲突、同结构名不一致等问题挡在编辑器外。

用法:
    python validate_tables.py <Tables目录 | 单个.xlsx>

行为约定与目标 Unity 项目内置的导表工具（ExportTablesTool.cs，
命名空间 KTXGame.ExportTables）及其随附的《导表工具使用说明.md》对齐：
    - Sheet 至少 4 行：第 1 行字段说明、第 2 行字段类型、第 3 行字段名、第 4 行起数据
    - 类型仅支持 int/int32/long/int64/float/single/double/bool/boolean/string
    - Sheet 名 "结构名.表名" 按第一个点拆分；最终表名全局唯一；
      同一结构名下所有 Sheet 的字段名/类型/顺序必须完全一致
    - 少于 4 行或无有效数据行的 Sheet 会被导表工具跳过

本脚本有意比导表工具更严格：工具对"字段名非法"是静默剥离/补字符（不报错），
最终在运行时才以反射失配的形式炸出来。脚本把这类问题提前判为 ERROR。
因此当脚本报 ERROR 而工具能跑通时，优先修 Excel，不要放宽校验规则。

退出码: 0 = 无 ERROR；1 = 存在 ERROR（WARNING 不影响退出码）。
本脚本只读取文件，不修改任何 Excel。
"""

import argparse
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("ERROR: 缺少依赖 openpyxl，请先在当前 Python 环境执行: pip install openpyxl")
    sys.exit(1)

SUPPORTED_TYPES = {
    "int": "int",
    "int32": "int",
    "long": "long",
    "int64": "long",
    "float": "float",
    "single": "float",
    "double": "double",
    "bool": "bool",
    "boolean": "bool",
    "string": "string",
}

# 允许出现在字段名中的字符：字母、数字、下划线
FIELD_NAME_OK = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_")


def cell_str(value):
    """把单元格值转成去空格字符串；None / 空白返回空串。"""
    if value is None:
        return ""
    return str(value).strip()


class Issue:
    def __init__(self, level, where, message, hint=""):
        self.level = level  # "ERROR" / "WARN"
        self.where = where
        self.message = message
        self.hint = hint


def trailing_blank(row):
    """判断一行是否全空。"""
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def check_data_value(value, col_type):
    """对数据行做宽松类型检查，返回警告文案；None 表示无问题。"""
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        if text.lower() == "null":
            return "数据单元格手填了文字 null（Excel 空白导出即为 null，不要手填）"
        if text == "":
            return None
        value = text
    else:
        text = value

    if col_type in ("int", "long"):
        if isinstance(value, bool):
            return f"整型列出现布尔值 {value}"
        if isinstance(value, int):
            return None
        if isinstance(value, float) and value.is_integer():
            return None
        if isinstance(value, str):
            try:
                int(text)
                return None
            except ValueError:
                return f"整型列出现无法解析为整数的文本: {value!r}"
        return f"整型列出现非整数值: {value!r}"

    if col_type in ("float", "double"):
        if isinstance(value, bool):
            return f"浮点列出现布尔值 {value}"
        if isinstance(value, (int, float)):
            return None
        if isinstance(value, str):
            try:
                float(text)
                return None
            except ValueError:
                return f"浮点列出现无法解析为数字的文本: {value!r}"
        return f"浮点列出现非数值: {value!r}"

    if col_type == "bool":
        if isinstance(value, bool):
            return None
        if isinstance(value, (int, float)) and value in (0, 1):
            return None
        if isinstance(value, str) and text.lower() in ("true", "false"):
            return None
        return f"布尔列出现无法识别的值: {value!r}（应为 true/false）"

    return None


def validate_sheet(wb_name, ws, issues, table_registry, schema_registry):
    where = f"{wb_name}!{ws.title}"
    sheet_name = ws.title.strip()

    # 收集所有行（去掉尾部全空行）
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    while rows and trailing_blank(rows[-1]):
        rows.pop()

    if len(rows) < 4:
        issues.append(Issue(
            "WARN", where,
            f"行数不足 4 行（当前 {len(rows)} 行），导表工具会跳过该 Sheet",
            "如果这是有效配置表，请补齐 4 行表头 + 数据行"))
        return

    header_desc, header_type, header_name = rows[0], rows[1], rows[2]

    # 表头行存在合并单元格会导致列信息缺失，明确提示
    merged_ranges = list(getattr(ws, "merged_cells", None) and ws.merged_cells.ranges or [])
    for mr in merged_ranges:
        if mr.min_row <= 3 and mr.max_row >= 1:
            issues.append(Issue(
                "ERROR", where,
                f"表头前 3 行存在合并单元格 {mr.coord}",
                "取消合并单元格，导表工具按逐列读取表头"))

    # 确定有效列范围：三行表头中任一非空即视为有效列
    col_count = max(len(header_desc), len(header_type), len(header_name))
    valid_cols = []
    for idx in range(col_count):
        d = cell_str(header_desc[idx]) if idx < len(header_desc) else ""
        t = cell_str(header_type[idx]) if idx < len(header_type) else ""
        n = cell_str(header_name[idx]) if idx < len(header_name) else ""
        if d or t or n:
            valid_cols.append((idx + 1, d, t, n))

    if not valid_cols:
        issues.append(Issue(
            "ERROR", where, "表头三行全空，没有可导出的列", "补齐第 1-3 行表头"))
        return

    # 逐列表头完整性 + 类型 + 字段名
    fields = []  # [(field_name, normalized_type, col_index)]
    for col_no, d, t, n in valid_cols:
        col_label = f"{where} 第{col_no}列"
        if not d or not t or not n:
            missing = [label for label, v in (("说明", d), ("类型", t), ("字段名", n)) if not v]
            issues.append(Issue(
                "ERROR", col_label,
                f"表头不完整，缺少: {'、'.join(missing)}",
                "每个有效列必须同时具备说明、类型、字段名三项"))
            continue

        t_lower = t.lower()
        if t_lower not in SUPPORTED_TYPES:
            issues.append(Issue(
                "ERROR", col_label,
                f"字段 {n} 的类型 {t!r} 不受支持",
                f"当前仅支持: {', '.join(sorted(SUPPORTED_TYPES))}"))
            continue

        normalized = SUPPORTED_TYPES[t_lower]
        if t != t_lower:
            issues.append(Issue(
                "WARN", col_label,
                f"类型 {t!r} 将被归一化为 {normalized}",
                "建议直接填写归一化后的小写类型"))

        if " " in n or any(ch not in FIELD_NAME_OK for ch in n):
            issues.append(Issue(
                "ERROR", col_label,
                f"字段名 {n!r} 含空格或非法字符",
                "字段名只用英文字母、数字、下划线"))
            continue

        # 工具生成 C# 字段时会对首字符为数字的名字补 '_'（1stItem -> _1stItem），
        # 但 JSON 键仍是原名，运行时 LoadData 按 JSON 键反射查找必然失败。
        # 工具本身不报错，属于静默失败，必须在这里拦住。
        if n[0].isdigit():
            issues.append(Issue(
                "ERROR", col_label,
                f"字段名 {n!r} 以数字开头",
                "导表工具会静默补下划线生成 C# 字段名（如 1stItem -> _1stItem），"
                "但 JSON 键保持原名，运行时反射查找必然失败且工具不报错；"
                "请改为以字母或下划线开头"))
            continue

        fields.append((n, normalized, col_no))

    # 同 Sheet 内字段名重复
    seen = {}
    for n, normalized, col_no in fields:
        if n in seen:
            issues.append(Issue(
                "ERROR", where,
                f"第 {col_no} 列字段名 {n!r} 与第 {seen[n]} 列重复",
                "字段名是 JSON 键名，必须唯一"))
        else:
            seen[n] = col_no

    # 解析 结构名.表名
    if "." in sheet_name:
        schema_name, table_name = sheet_name.split(".", 1)
        schema_name, table_name = schema_name.strip(), table_name.strip()
        if not schema_name or not table_name:
            issues.append(Issue(
                "ERROR", where,
                f"Sheet 名 {sheet_name!r} 的 结构名.表名 拆分结果为空",
                "按 结构名.表名 命名，且两段都不能为空"))
            return
    else:
        schema_name = table_name = sheet_name

    # 最终表名全局唯一
    if table_name.lower() in table_registry:
        issues.append(Issue(
            "ERROR", where,
            f"最终表名 {table_name!r} 与 {table_registry[table_name.lower()]} 重复",
            "tableName 必须全局唯一（含 结构名.表名 拆分后的表名）"))
    else:
        table_registry[table_name.lower()] = where

    # 数据行检查（宽松：只报 WARNING，工具实际行为以导出为准）
    data_rows = rows[3:]
    valid_data_rows = [r for r in data_rows if not trailing_blank(r)]
    if not valid_data_rows:
        issues.append(Issue(
            "WARN", where, "没有任何有效数据行，导表工具会跳过该 Sheet",
            "如果这是有效配置表，请填入数据"))
    else:
        for row_no, row in enumerate(data_rows, start=4):
            if trailing_blank(row):
                continue
            for n, normalized, col_no in fields:
                value = row[col_no - 1] if col_no <= len(row) else None
                warn = check_data_value(value, normalized)
                if warn:
                    issues.append(Issue(
                        "WARN", f"{where} 第{row_no}行第{col_no}列", warn, ""))
                    break  # 同一行只报一条，避免刷屏

    # 登记 schema 用于一致性检查
    schema_registry.setdefault(schema_name, []).append(
        (where, [(n, t) for n, t, _ in fields]))


def validate_workbook(xlsx_path, issues, table_registry, schema_registry):
    try:
        wb = load_workbook(xlsx_path, data_only=True, read_only=False)
    except Exception as exc:  # 损坏文件、Office 锁定等
        issues.append(Issue(
            "ERROR", str(xlsx_path), f"无法打开工作簿: {exc}", "确认文件未被 Office 锁定且未损坏"))
        return
    try:
        for ws in wb.worksheets:
            validate_sheet(xlsx_path.name, ws, issues, table_registry, schema_registry)
    finally:
        wb.close()


def collect_xlsx(target: Path):
    if target.is_file():
        return [target] if target.suffix.lower() == ".xlsx" and not target.name.startswith("~$") else []
    files = []
    for p in sorted(target.rglob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        files.append(p)
    return files


def main():
    parser = argparse.ArgumentParser(
        description="Excel 配置表导表前置校验（只读，不改 Excel）")
    parser.add_argument("path", help="Tables 目录（递归）或单个 .xlsx 文件")
    args = parser.parse_args()

    target = Path(args.path)
    if not target.exists():
        print(f"ERROR: 路径不存在: {target}")
        return 1

    files = collect_xlsx(target)
    if not files:
        print(f"未找到可校验的 .xlsx 文件: {target}")
        return 1

    issues = []
    table_registry = {}
    schema_registry = {}

    for f in files:
        validate_workbook(f, issues, table_registry, schema_registry)

    # 同结构名一致性检查：字段名、类型、顺序必须完全一致（字段说明可不同）
    for schema_name, members in sorted(schema_registry.items()):
        if len(members) < 2:
            continue
        base_where, base_fields = members[0]
        for other_where, other_fields in members[1:]:
            if other_fields != base_fields:
                base_names = [n for n, _ in base_fields]
                other_names = [n for n, _ in other_fields]
                issues.append(Issue(
                    "ERROR", other_where,
                    f"结构名 {schema_name!r} 下的字段结构与其成员不一致",
                    f"以 {base_where}（字段: {', '.join(base_names)}）为基准；"
                    f"当前（字段: {', '.join(other_names)}）。"
                    f"同一结构名下所有 Sheet 的字段名、类型、顺序必须完全一致，"
                    f"如行为确实不同请改用不同结构名"))

    # 输出报告
    errors = [i for i in issues if i.level == "ERROR"]
    warns = [i for i in issues if i.level == "WARN"]

    print(f"校验文件数: {len(files)}")
    for f in files:
        print(f"  - {f}")
    print()

    if not issues:
        print("结果: 全部通过，无 ERROR / WARNING")
        print("可以进入 Unity 执行 Tools/Export Tables。")
        return 0

    for i in errors:
        print(f"[ERROR] {i.where}: {i.message}")
        if i.hint:
            print(f"        修复方向: {i.hint}")
    for i in warns:
        print(f"[WARN ] {i.where}: {i.message}")
        if i.hint:
            print(f"        说明: {i.hint}")

    print()
    print(f"结果: {len(errors)} 个 ERROR, {len(warns)} 个 WARNING")
    if errors:
        print("存在 ERROR：先修复 Excel 再进 Unity 导出。")
        return 1
    print("无 ERROR，WARNING 需人工判断（多为非阻断项）。可以进入 Unity 导出。")
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    sys.exit(main())
